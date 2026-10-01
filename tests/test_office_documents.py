"""Word + Excel handling: building .xlsx deliverables, reading .docx/.xlsx attachments, creating/editing drafts.

Everything here is a draft for review - nothing is ever emailed by these paths (email_send stays approval-gated)."""

import base64
import io
import zipfile

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from jarvis.core import Jarvis
from jarvis.integrations.microsoft365 import select_office_attachments
from jarvis.main import create_app
from jarvis.services import documents
from tests.fakes import FakeClient

GOOD_ID = "b" * 32

STOCK = """## Stock levels

| Item | Qty | Unit cost |
|---|---|---|
| Optical detector | 42 | £12.50 |
| Sounder base | 1,200 | 3.1 |
| Postcode | 0123 | =HYPERLINK("http://evil.example","x") |

## Notes

Reorder anything under **10**.
"""


def make_doc(markdown=STOCK, **over):
    return {"id": GOOD_ID, "created_at": "2026-09-30T10:00:00+00:00", "kind": "stock_export",
            "title": "Stock export", "markdown": markdown, **over}


def _jarvis(tmp_path):
    from jarvis.config import Settings

    return Jarvis(Settings(data_dir=tmp_path, scheduler_enabled=False, _env_file=None), client=FakeClient())


def test_render_xlsx_builds_a_sheet_per_table():
    data = documents.render_xlsx(make_doc(), "Salts Fire and Security")
    assert data.startswith(b"PK")
    wb = load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["Stock levels", "Notes"]
    ws = wb["Stock levels"]
    assert [c.value for c in ws[1]] == ["Item", "Qty", "Unit cost"]
    assert ws["B2"].value == 42 and ws["C2"].value == 12.5  # numbers become real numbers (£ stripped)
    assert ws["B3"].value == 1200 and ws["C3"].value == 3.1
    assert ws["B4"].value == "0123"  # leading zeros are kept as text
    assert ws.freeze_panes == "A2"
    assert wb["Notes"]["A1"].value == "Reorder anything under 10."  # inline markdown stripped
    assert wb.properties.title == "Stock export"


def test_render_xlsx_never_creates_formulas():
    wb = load_workbook(io.BytesIO(documents.render_xlsx(make_doc(), "Salts")))
    cell = wb["Stock levels"]["C4"]
    assert cell.data_type == "s" and cell.value.startswith("=HYPERLINK")
    with zipfile.ZipFile(io.BytesIO(documents.render_xlsx(make_doc(), "Salts"))) as z:
        sheet_xml = z.read("xl/worksheets/sheet1.xml").decode()
    assert "<f>" not in sheet_xml


def test_render_xlsx_without_tables_and_odd_titles():
    wb = load_workbook(io.BytesIO(documents.render_xlsx(make_doc("# Report\n\nJust text\x07 here.\n\n- a\n- b"), "S")))
    assert wb.sheetnames == ["Document"]
    values = [r[0].value for r in wb["Document"].iter_rows()]
    assert "Just text here." in values and "a" in values
    # sheet names are sanitised (Excel forbids []:*?/\ and > 31 chars) and made unique
    md = "## A/B:C*" + "x" * 50 + "\n\n| a |\n|---|\n| 1 |\n\n## A/B:C*" + "x" * 50 + "\n\n| a |\n|---|\n| 2 |\n"
    wb = load_workbook(io.BytesIO(documents.render_xlsx(make_doc(md), "S")))
    assert len(wb.sheetnames) == 2 and len(set(wb.sheetnames)) == 2
    assert all(len(n) <= 31 and not set("[]:*?/\\") & set(n) for n in wb.sheetnames)
    assert load_workbook(io.BytesIO(documents.render_xlsx(make_doc(""), "S"))).sheetnames == ["Document"]


def _xlsx_bytes():
    wb = Workbook()
    ws = wb.active
    ws.title = "Jobs"
    ws.append(["Ref", "Site", "Hours", "Date"])
    ws.append(["J-1", "Aire Valley | Care Home", 2.0, None])
    ws.append([None, None, None, None])
    ws.append(["J-2", "Mill", 3.5, None])
    ws2 = wb.create_sheet("Empty")
    ws2["A1"] = None
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_xlsx_to_markdown_reads_sheets_and_tables():
    md = documents.xlsx_to_markdown(_xlsx_bytes())
    assert "## Jobs" in md and "| Ref | Site | Hours | Date |" in md
    assert "| J-1 | Aire Valley / Care Home | 2 |  |" in md  # integral floats tidy, pipes can't break the table
    assert "| J-2 | Mill | 3.5 |  |" in md
    assert "## Empty" in md and "(empty sheet)" in md
    assert [b[0] for b in documents._parse_markdown(md)][:2] == ["h", "table"]  # round-trips through our parser


def test_xlsx_to_markdown_truncates_big_sheets():
    wb = Workbook()
    ws = wb.active
    for i in range(documents.MAX_SHEET_ROWS + 50):
        ws.append([f"row{i}"])
    buf = io.BytesIO()
    wb.save(buf)
    md = documents.xlsx_to_markdown(buf.getvalue())
    assert f"row{documents.MAX_SHEET_ROWS - 1}" in md and f"row{documents.MAX_SHEET_ROWS + 1}" not in md
    assert "truncated" in md


def test_docx_round_trip_through_markdown():
    docx = documents.render_docx(make_doc(markdown="# Title one\n\nHello **world**.\n\n| A | B |\n|---|---|\n| 1 | 2 |\n\n"
                                                   "- point one\n- point two\n\n1. first\n"), "Salts")
    md = documents.docx_to_markdown(docx)
    assert "# Title one" in md and "Hello world." in md
    assert "| A | B |" in md and "| 1 | 2 |" in md
    assert "- point one" in md and "1. first" in md
    assert "PROTECTING WHAT MATTERS MOST" in md  # the title page text is read too


def test_readers_reject_bad_files(monkeypatch):
    with pytest.raises(ValueError):
        documents.xlsx_to_markdown(b"not a zip")
    with pytest.raises(ValueError):
        documents.docx_to_markdown(b"not a zip")
    with pytest.raises(ValueError):
        documents.office_to_markdown("notes.txt", b"hello")
    # zip bomb guard: declared uncompressed size beyond the cap
    monkeypatch.setattr(documents, "MAX_UNZIPPED_BYTES", 1000)
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("word/document.xml", b"0" * 2000)
    with pytest.raises(ValueError):
        documents.docx_to_markdown(buf.getvalue())
    with pytest.raises(ValueError):
        documents.xlsx_to_markdown(buf.getvalue())


def test_office_to_markdown_dispatches_on_extension():
    assert "## Jobs" in documents.office_to_markdown("Jobs.XLSX", _xlsx_bytes())
    assert "Heading" in documents.office_to_markdown("a.docx", documents.render_docx(make_doc("# Heading"), "S"))


def test_select_office_attachments_filters_and_caps():
    def att(name, size=100, kind="#microsoft.graph.fileAttachment", content="QQ=="):
        return {"@odata.type": kind, "name": name, "size": size, "contentBytes": content,
                "contentType": "application/octet-stream"}

    items = [att("Tender.DOCX"), att("stock.xlsx"), att("report.pdf"), att("macro.xlsm"), att("big.docx", 101),
             att("ref.docx", kind="#microsoft.graph.referenceAttachment"), att("empty.xlsx", content="")]
    picked = select_office_attachments(items, max_bytes=100)
    assert [a["name"] for a in picked] == ["Tender.DOCX", "stock.xlsx"]
    assert picked[0] == {"name": "Tender.DOCX", "data": "QQ=="}


async def test_read_attachment_tool_returns_text_and_flags_it_untrusted(tmp_path):
    from jarvis.brain.tools import TOOLS_BY_NAME, AttachmentReadIn, email_attachment_read

    assert TOOLS_BY_NAME["email_attachment_read"].approval is False
    j = _jarvis(tmp_path)
    seen = {}

    async def office(message_id, mailbox=None, max_bytes=0):
        seen["id"] = message_id
        return [{"name": "jobs.xlsx", "data": base64.b64encode(_xlsx_bytes()).decode()},
                {"name": "bad.docx", "data": base64.b64encode(b"junk").decode()}]

    j.mail.office_attachments = office
    out = await email_attachment_read(j, AttachmentReadIn(message_id="demo-1"))
    assert seen["id"] == "demo-1"
    assert out["attachments"][0]["name"] == "jobs.xlsx" and "| J-2 | Mill | 3.5 |" in out["attachments"][0]["text"]
    assert out["attachments"][1]["name"] == "bad.docx" and "couldn't read" in out["attachments"][1]["error"]
    assert "untrusted" in out["note"].lower()
    only = await email_attachment_read(j, AttachmentReadIn(message_id="demo-1", name="JOBS.xlsx"))
    assert [a["name"] for a in only["attachments"]] == ["jobs.xlsx"]
    await j.http.aclose()


async def test_read_attachment_tool_with_nothing_attached(tmp_path):
    from jarvis.brain.tools import AttachmentReadIn, email_attachment_read

    j = _jarvis(tmp_path)
    out = await email_attachment_read(j, AttachmentReadIn(message_id="demo-1"))
    assert out["attachments"] == [] and "no Word or Excel" in out["note"]
    await j.http.aclose()


async def test_create_document_is_a_stored_draft_and_never_sends(tmp_path):
    from jarvis.brain.tools import TOOLS_BY_NAME, OfficeDocumentIn, draft_office_document

    tool = TOOLS_BY_NAME["draft_office_document"]
    assert tool.approval is False and "email_send" in tool.description and "never sent" in tool.description
    j = _jarvis(tmp_path)

    async def boom(*a, **k):
        raise AssertionError("an email was sent")

    j.mail.send_mail = boom
    j.mail.create_reply_draft = boom
    q = j.bus.subscribe()
    out = await draft_office_document(j, OfficeDocumentIn(format="xlsx", kind="stock_export", title="Van stock",
                                                         content=STOCK))
    assert documents.valid_doc_id(out["doc_id"])
    assert out["download_url"] == f"/api/documents/{out['doc_id']}/xlsx" and "not been sent" in out["note"]
    stored = j.documents.get(out["doc_id"])
    assert stored["kind"] == "stock_export" and stored["title"] == "Van stock" and stored["markdown"] == STOCK
    events = []
    while not q.empty():
        events.append(q.get_nowait())
    assert any(e["type"] == "display" and e["data"].get("doc_id") == out["doc_id"] for e in events)
    out = await draft_office_document(j, OfficeDocumentIn(format="docx", title="Tender", content="# Hi"))
    assert out["download_url"].endswith("/docx") and j.documents.get(out["doc_id"])["kind"] == "report"
    await j.http.aclose()


async def test_edit_attachment_makes_a_new_draft_from_the_original(tmp_path):
    from jarvis.brain.tools import TOOLS_BY_NAME, OfficeEditIn, edit_office_document

    assert TOOLS_BY_NAME["edit_office_document"].approval is False
    j = _jarvis(tmp_path)

    async def office(message_id, mailbox=None, max_bytes=0):
        return [{"name": "jobs.xlsx", "data": base64.b64encode(_xlsx_bytes()).decode()}]

    j.mail.office_attachments = office
    j.client.beta.messages.default_text = "## Jobs\n\n| Ref | Site |\n|---|---|\n| J-1 | Changed |"
    out = await edit_office_document(j, OfficeEditIn(instructions="Rename the site", format="xlsx",
                                                     message_id="demo-1", attachment_name="jobs.xlsx"))
    sent = j.client.beta.messages.calls[-1]["messages"][0]["content"]
    assert "Rename the site" in sent and "| J-2 | Mill | 3.5 |" in sent and "<original>" in sent
    stored = j.documents.get(out["doc_id"])
    assert "Changed" in stored["markdown"] and stored["title"] == "Edited - jobs.xlsx"
    assert "original is unchanged" in out["note"] and "formulas" in out["note"]
    await j.http.aclose()


async def test_edit_stored_draft_and_error_paths(tmp_path):
    from jarvis.brain.tools import OfficeEditIn, edit_office_document

    j = _jarvis(tmp_path)
    j.db.add_document(GOOD_ID, "report", "Monthly report", "# Report\n\nOld text")
    out = await edit_office_document(j, OfficeEditIn(instructions="Update it", format="docx", doc_id=GOOD_ID))
    assert j.documents.get(out["doc_id"])["title"] == "Edited - Monthly report" and out["doc_id"] != GOOD_ID
    assert j.documents.get(GOOD_ID)["markdown"] == "# Report\n\nOld text"  # original draft untouched
    before = len(j.client.beta.messages.calls)
    for bad in (OfficeEditIn(instructions="x", format="docx"),
                OfficeEditIn(instructions="x", format="docx", doc_id="c" * 32),
                OfficeEditIn(instructions="x", format="docx", message_id="demo-1", attachment_name="nope.docx")):
        res = await edit_office_document(j, bad)
        assert "error" in res and "doc_id" not in res
    assert len(j.client.beta.messages.calls) == before  # nothing sent to the model when there was no source
    await j.http.aclose()


def test_download_endpoint_serves_xlsx(settings):
    settings.jarvis_owner_password = "s3cret"
    j = Jarvis(settings, client=FakeClient())
    j.db.add_document(GOOD_ID, "stock_export", "Stock / export", STOCK)
    app = create_app(settings, j)
    with TestClient(app) as c:
        assert c.get(f"/api/documents/{GOOD_ID}/xlsx").status_code == 401  # not signed in
        c.post("/login", data={"password": "s3cret"}, follow_redirects=False)
        r = c.get(f"/api/documents/{GOOD_ID}/xlsx")
        assert r.status_code == 200 and r.headers["content-type"] == documents.XLSX_MIME
        assert r.headers["content-disposition"] == 'attachment; filename="Stock-export.xlsx"'
        assert load_workbook(io.BytesIO(r.content)).sheetnames[0] == "Stock levels"
        assert c.get(f"/api/documents/{GOOD_ID}/exe").status_code == 404
